"""내보낸 PPTX 에 날짜·쪽번호 자리가 남아 있지 않은가.

python-pptx 의 기본 서식은 모든 레이아웃과 마스터에 날짜·바닥글·쪽번호
자리를 넣어 둔다. 우리 덱은 쪽번호와 출처를 글로 직접 그리므로 그 자리는
쓰이지 않는데, 파워포인트로 열었더니 그린 쪽번호 옆에 날짜와 쪽번호가 한 번
더 보였다 (2026-09-16). `_strip_base_footer_placeholders` 가 그 자리를 지우고
머리글·바닥글을 끈다. 이 검사는 그 약속 하나만 본다.

    python3 -m unittest tests.test_footer_placeholders
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent / ".claude/skills/ppt-master/scripts"
sys.path.insert(0, str(SCRIPTS))

from svg_to_pptx.pptx_package import builder  # noqa: E402

FOOTER_TYPES = ('type="dt"', 'type="ftr"', 'type="sldNum"')


class StockFooterPlaceholders(unittest.TestCase):
    def _extracted_stock_package(self, tmp: Path) -> Path:
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        base = tmp / "base.pptx"
        prs.save(str(base))
        out = tmp / "x"
        with zipfile.ZipFile(base) as zf:
            zf.extractall(out)
        return out

    def test_기본_서식의_날짜_바닥글_쪽번호_자리를_지운다(self):
        with tempfile.TemporaryDirectory() as d:
            root = self._extracted_stock_package(Path(d))
            parts = list(root.glob("ppt/slideMasters/*.xml")) + list(root.glob("ppt/slideLayouts/*.xml"))
            before = sum(any(t in p.read_text(encoding="utf-8") for t in FOOTER_TYPES) for p in parts)
            self.assertGreater(before, 0, "기본 서식에 자리가 없다면 이 검사는 뜻이 없다")

            removed = builder._strip_base_footer_placeholders(root)

            self.assertGreater(removed, 0)
            for p in parts:
                text = p.read_text(encoding="utf-8")
                for t in FOOTER_TYPES:
                    self.assertNotIn(t, text, f"{p.name} 에 {t} 자리가 남았다")
                self.assertIn('sldNum="0"', text, f"{p.name} 의 머리글·바닥글이 꺼지지 않았다")

    def test_두_번_돌려도_같다(self):
        with tempfile.TemporaryDirectory() as d:
            root = self._extracted_stock_package(Path(d))
            builder._strip_base_footer_placeholders(root)
            self.assertEqual(builder._strip_base_footer_placeholders(root), 0)


if __name__ == "__main__":
    unittest.main()
